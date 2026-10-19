
import json
import random
import textwrap
from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer
)
from xml.sax.saxutils import escape


BASE_DIR = Path(__file__).resolve().parent
RESUME_DIR = BASE_DIR / "resumes"

SEED = 42

ROLE_DATA = {
    "Python Backend Developer": {
        "core": ["Python", "FastAPI", "REST APIs"],
        "additional": [
            "Django", "PostgreSQL", "Docker",
            "Redis", "AWS", "SQL", "Git"
        ],
        "projects": [
            "Built a REST API for an inventory platform",
            "Developed asynchronous order-processing services",
            "Implemented authentication and database APIs"
        ]
    },
    "Frontend Developer": {
        "core": ["JavaScript", "React", "HTML", "CSS"],
        "additional": [
            "TypeScript", "Angular", "Redux",
            "Jest", "Git", "REST APIs"
        ],
        "projects": [
            "Built a responsive analytics dashboard",
            "Developed reusable frontend components",
            "Improved accessibility in a web application"
        ]
    },
    "Full Stack Developer": {
        "core": ["Python", "React", "SQL"],
        "additional": [
            "FastAPI", "Django", "PostgreSQL",
            "Docker", "JavaScript", "AWS", "Git"
        ],
        "projects": [
            "Built a full-stack task management application",
            "Implemented frontend and backend integrations",
            "Developed an internal project tracking platform"
        ]
    },
    "Machine Learning Engineer": {
        "core": [
            "Python", "Machine Learning", "Scikit-learn"
        ],
        "additional": [
            "PyTorch", "TensorFlow", "Pandas",
            "NumPy", "Docker", "SQL", "AWS"
        ],
        "projects": [
            "Developed a customer churn prediction model",
            "Built a text classification pipeline",
            "Evaluated a recommendation model"
        ]
    },
    "DevOps Engineer": {
        "core": ["Docker", "Kubernetes", "AWS"],
        "additional": [
            "Terraform", "Jenkins", "Linux",
            "Git", "Python", "CI/CD", "Prometheus"
        ],
        "projects": [
            "Automated cloud infrastructure deployments",
            "Created container-based deployment pipelines",
            "Implemented monitoring for distributed services"
        ]
    },
    "Data Engineer": {
        "core": ["Python", "SQL", "ETL"],
        "additional": [
            "Apache Spark", "Airflow", "PostgreSQL",
            "AWS", "Docker", "Pandas", "Kafka"
        ],
        "projects": [
            "Built a batch data ingestion pipeline",
            "Created an ETL workflow for analytics",
            "Optimized SQL transformations for reporting"
        ]
    },
    "QA Automation Engineer": {
        "core": ["Python", "Selenium", "Testing"],
        "additional": [
            "Pytest", "Postman", "REST APIs",
            "Jenkins", "Git", "SQL", "Docker"
        ],
        "projects": [
            "Developed automated regression tests",
            "Built an API testing framework",
            "Integrated automated tests into CI/CD"
        ]
    }
}

FIRST_NAMES = [
    "Aarav", "Ananya", "Rohan", "Priya", "Kabir",
    "Meera", "Aditya", "Sneha", "Arjun", "Isha",
    "Neha", "Vikram", "Tara", "Nikhil", "Diya"
]

LAST_NAMES = [
    "Sharma", "Verma", "Gupta", "Rao", "Kapoor",
    "Mehta", "Singh", "Nair", "Patel", "Joshi"
]

EDUCATION = [
    "B.Tech in Computer Science",
    "B.E. in Information Technology",
    "B.Sc. in Computer Science",
    "MCA",
    "M.Tech in Computer Science"
]

EXPERIENCE_LEVELS = [2, 4, 6, 8, 10]

JOB_SPECS = [
    {
        "job_id": "JD001",
        "title": "Python Backend Developer",
        "must_have_skills": ["Python", "FastAPI"],
        "preferred_skills": ["PostgreSQL", "Docker", "AWS"],
        "min_experience": 5
    },
    {
        "job_id": "JD002",
        "title": "Frontend Developer",
        "must_have_skills": ["React", "JavaScript"],
        "preferred_skills": ["TypeScript", "Angular", "Jest"],
        "min_experience": 3
    },
    {
        "job_id": "JD003",
        "title": "Full Stack Developer",
        "must_have_skills": ["Python", "React", "SQL"],
        "preferred_skills": ["FastAPI", "Docker", "AWS"],
        "min_experience": 4
    },
    {
        "job_id": "JD004",
        "title": "Machine Learning Engineer",
        "must_have_skills": ["Python", "Machine Learning"],
        "preferred_skills": ["PyTorch", "TensorFlow", "AWS"],
        "min_experience": 3
    },
    {
        "job_id": "JD005",
        "title": "DevOps Engineer",
        "must_have_skills": ["Docker", "Kubernetes", "AWS"],
        "preferred_skills": ["Terraform", "Jenkins", "Linux"],
        "min_experience": 4
    }
]


def save_json(path, data):
    path.write_text(
        json.dumps(data, indent=4, ensure_ascii=False),
        encoding="utf-8"
    )


def make_candidate(candidate_id, role, level, rng):
    role_info = ROLE_DATA[role]

    first = FIRST_NAMES[(candidate_id - 1) % len(FIRST_NAMES)]
    last = LAST_NAMES[(candidate_id * 3) % len(LAST_NAMES)]

    # Some profiles intentionally lack one core skill.
    core = list(role_info["core"])

    if level == 0:
        core = core[:-1]

    additional = rng.sample(
        role_info["additional"],
        k=min(3 + level // 2, len(role_info["additional"]))
    )

    skills = list(dict.fromkeys(core + additional))

    # Ensure a mixture of strong and borderline matches.
    # Do not infer skills from job labels.
    experience = EXPERIENCE_LEVELS[level]

    project = role_info["projects"][
        level % len(role_info["projects"])
    ]

    return {
        "candidate_id": f"C{candidate_id:03d}",
        "candidate_name": f"{first} {last}",
        "role": role,
        "experience_years": experience,
        "skills": skills,
        "education": rng.choice(EDUCATION),
        "project": project,
        "resume_path": ""
    }


def resume_text(candidate):
    skills = ", ".join(candidate["skills"])
    years = candidate["experience_years"]

    return (
        f"{candidate['candidate_name']}\n"
        f"{candidate['role']}\n\n"
        "PROFESSIONAL SUMMARY\n"
        f"{candidate['role']} with {years} years of experience "
        "delivering software projects and collaborating "
        "with cross-functional teams.\n\n"
        "SKILLS\n"
        f"{skills}\n\n"
        "EXPERIENCE\n"
        f"{years} years of professional experience.\n"
        f"Worked as a {candidate['role']} on software "
        "development and delivery assignments.\n"
        f"Applied the following technologies: {skills}.\n\n"
        "PROJECTS\n"
        f"{candidate['project']}.\n"
        "Responsibilities included implementation, "
        "testing, documentation, and code reviews.\n\n"
        "EDUCATION\n"
        f"{candidate['education']}\n"
    )


def write_txt(path, content):
    path.write_text(content, encoding="utf-8")


def write_docx(path, content):
    document = Document()

    lines = content.splitlines()

    for index, line in enumerate(lines):
        if not line.strip():
            continue

        if index == 0:
            document.add_heading(line, level=0)
        elif line in {
            "PROFESSIONAL SUMMARY",
            "SKILLS",
            "EXPERIENCE",
            "PROJECTS",
            "EDUCATION"
        }:
            document.add_heading(line.title(), level=1)
        else:
            document.add_paragraph(line)

    document.save(path)


def write_pdf(path, content):
    styles = getSampleStyleSheet()

    document = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=48,
        rightMargin=48,
        topMargin=48,
        bottomMargin=48
    )

    story = []

    for index, line in enumerate(content.splitlines()):
        if not line.strip():
            story.append(Spacer(1, 8))
            continue

        if index == 0:
            style = styles["Title"]
        elif line in {
            "PROFESSIONAL SUMMARY",
            "SKILLS",
            "EXPERIENCE",
            "PROJECTS",
            "EDUCATION"
        }:
            style = styles["Heading2"]
        else:
            style = styles["BodyText"]

        # Escape XML characters before ReportLab Paragraph.
        for wrapped_line in textwrap.wrap(
            line, width=100, break_long_words=False
        ) or [""]:
            story.append(
                Paragraph(escape(wrapped_line), style)
            )

    document.build(story)


def build_job_description(spec):
    mandatory = ", ".join(spec["must_have_skills"])
    preferred = ", ".join(spec["preferred_skills"])

    return (
        f"We are hiring a {spec['title']} with at least "
        f"{spec['min_experience']} years of experience. "
        f"Mandatory technical skills: {mandatory}. "
        f"Preferred skills include: {preferred}. "
        "The candidate will collaborate with engineering "
        "teams to deliver reliable software solutions."
    )


def label_candidate(candidate, job):
    """
    Ground truth is defined using structured attributes.

    Relevant means:
    - All mandatory skills are present.
    - Minimum experience is satisfied.

    Preferred skills are not mandatory.
    """
    candidate_skills = {
        skill.casefold() for skill in candidate["skills"]
    }

    missing = [
        skill for skill in job["must_have_skills"]
        if skill.casefold() not in candidate_skills
    ]

    experience_ok = (
        candidate["experience_years"]
        >= job["min_experience"]
    )

    relevant = not missing and experience_ok

    reasons = []

    if missing:
        reasons.append(
            "Missing mandatory skills: " + ", ".join(missing)
        )

    if not experience_ok:
        reasons.append(
            "Insufficient experience"
        )

    if relevant:
        reasons.append(
            "All mandatory skills and minimum "
            "experience requirements satisfied"
        )

    return {
        "candidate_id": candidate["candidate_id"],
        "resume_path": candidate["resume_path"],
        "relevant": relevant,
        "relevance_grade": 2 if relevant else 0,
        "reasons": reasons
    }


def generate_dataset(output_dir=BASE_DIR, seed=SEED):
    output_dir = Path(output_dir).resolve()
    resume_dir = output_dir / "resumes"
    resume_dir.mkdir(parents=True, exist_ok=True)

    rng = random.Random(seed)

    candidates = []
    candidate_number = 0

    for role in ROLE_DATA:
        for level in range(5):
            candidate_number += 1

            candidate = make_candidate(
                candidate_number,
                role,
                level,
                rng
            )

            # Exactly 15 PDF, 10 DOCX, and 10 TXT files.
            if candidate_number <= 15:
                extension = ".pdf"
            elif candidate_number <= 25:
                extension = ".docx"
            else:
                extension = ".txt"

            filename = (
                f"{candidate['candidate_id']}_"
                f"{role.lower().replace(' ', '_')}"
                f"{extension}"
            )

            candidate["resume_path"] = (
                f"resumes/{filename}"
            )

            destination = resume_dir / filename
            content = resume_text(candidate)

            if extension == ".pdf":
                write_pdf(destination, content)
            elif extension == ".docx":
                write_docx(destination, content)
            else:
                write_txt(destination, content)

            candidates.append(candidate)

    jobs = []

    for spec in JOB_SPECS:
        job = dict(spec)
        job["job_description"] = build_job_description(spec)
        jobs.append(job)

    relevance = {}

    for job in jobs:
        relevance[job["job_id"]] = {
            "job_title": job["title"],
            "judgments": [
                label_candidate(candidate, job)
                for candidate in candidates
            ]
        }

    save_json(output_dir / "candidates.json", candidates)
    save_json(output_dir / "job_descriptions.json", jobs)
    save_json(output_dir / "relevance_labels.json", relevance)

    return {
        "candidates": candidates,
        "jobs": jobs,
        "relevance": relevance,
        "output_dir": str(output_dir)
    }


if __name__ == "__main__":
    result = generate_dataset()

    print("Dataset generated successfully!")
    print("Resumes:", len(result["candidates"]))
    print("Job descriptions:", len(result["jobs"]))

    for job_id, data in result["relevance"].items():
        relevant_count = sum(
            judgment["relevant"]
            for judgment in data["judgments"]
        )

        print(
            f"{job_id}: {relevant_count} relevant resumes"
        )

    print("Output directory:", result["output_dir"])
